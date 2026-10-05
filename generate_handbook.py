import os
import subprocess

def install_and_import(package):
    try:
        __import__(package)
    except ImportError:
        subprocess.check_call(["python", "-m", "pip", "install", package])

install_and_import("docx")

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.add_heading('SSP Fuel Lid Inspection System - Project Handbook', 0)

# 1. Project Overview
doc.add_heading('1. Project Overview', level=1)
doc.add_paragraph('The SSP Fuel Lid Inspection system is a production-ready, automated quality assurance solution. It utilizes computer vision and artificial intelligence (YOLOv8 & PaddleOCR) to inspect automotive fuel lids (doors) for defects and accurately read serial numbers. The system is designed to run autonomously on a Windows Edge PC, processing images from industrial cameras (GigE) or webcams, and providing real-time feedback via a web interface.')

# 2. Key Features
doc.add_heading('2. Key Features', level=1)
features = doc.add_paragraph()
features.add_run('1. AI-Powered Defect Detection: ').bold = True
features.add_run('Uses YOLOv8 to identify various surface defects such as dents, bulges, line marks, and damages on both the front and back of the fuel lids.\n')
features.add_run('2. Optical Character Recognition (OCR): ').bold = True
features.add_run('Employs PaddleOCR to read the serial numbers printed on the back panel of the fuel lids.\n')
features.add_run('3. Real-Time Web Interface: ').bold = True
features.add_run('Provides a real-time monitoring UI hosted via Flask, displaying the camera feed, current detections, and cycle status (e.g., waiting for part, front captured, back captured).\n')
features.add_run('4. Automated Reporting: ').bold = True
features.add_run('Generates professional A4 PDF inspection reports for every processed part using ReportLab.\n')
features.add_run('5. Smart Part Presence Detection: ').bold = True
features.add_run('Uses HSV color space masks (filtering out blue foam trays and detecting metallic surfaces) to accurately determine if a part is present.\n')
features.add_run('6. Configurable via YAML: ').bold = True
features.add_run('All critical settings (thresholds, timeouts, retention days, camera settings) are easily adjustable in config.yaml without altering the codebase.\n')
features.add_run('7. Data Retention & Cleanup: ').bold = True
features.add_run('A background thread automatically deletes old images and reports (configurable, default 30 days) to prevent disk space exhaustion.\n')
features.add_run('8. Windows Service Integration: ').bold = True
features.add_run('Can be installed as a background Windows service using NSSM, ensuring automatic startup and crash recovery.')

# 3. What Else Can It Do?
doc.add_heading('3. What Else Can It Do?', level=1)
doc.add_paragraph('Beyond simple inspection, the system can:')
capabilities = doc.add_paragraph(style='List Bullet')
capabilities = doc.add_paragraph('Support both standard and circular fuel door types automatically.')
capabilities = doc.add_paragraph('Rotate and dynamically analyze OCR crops to find text at various orientations.')
capabilities = doc.add_paragraph('Maintain separate AI confidence thresholds for different defect classes to fine-tune accuracy.')
capabilities = doc.add_paragraph('Rotate and manage logs automatically, keeping long-term operations stable.')
capabilities = doc.add_paragraph('Integrate with industrial IKapC SDK for specialized vision cameras (GigE), and standard RTSP IP cameras.')

# 4. Problems & Solutions
doc.add_heading('4. Problems & Solutions', level=1)
doc.add_heading('Problem 1: Disk Space Exhaustion', level=2)
doc.add_paragraph('Solution: The system saves high-resolution images and PDFs constantly. To solve disk space issues, an auto-cleanup background thread deletes data older than the specified retention period (30 days).')

doc.add_heading('Problem 2: Hardcoded Parameters Blocking Tuning', level=2)
doc.add_paragraph('Solution: Previous versions required developer intervention to change AI thresholds. Now, a `config.yaml` file allows on-site operators to tune confidence thresholds, timings, and storage limits directly.')

doc.add_heading('Problem 3: False Positives on Empty Trays', level=2)
doc.add_paragraph('Solution: The AI sometimes hallucinated defects on the blue foam tray. The solution implemented uses OpenCV HSV masking to detect the presence of the metallic gray part and explicitly filter out the blue background.')

doc.add_heading('Problem 4: Application Crashes Stopping Production', level=2)
doc.add_paragraph('Solution: Included `install_service.bat` uses NSSM to turn the application into a robust Windows service. If it crashes, it restarts automatically. It also starts on boot.')

doc.add_heading('Problem 5: Thread Contention between YOLO and PaddleOCR', level=2)
doc.add_paragraph('Solution: Explicitly limits environment variables (`OMP_NUM_THREADS`, etc.) to 1 and offloads PaddleOCR to a separate microservice (`paddleocr_server.py`) to prevent CPU thrashing and ensure stable FPS.')

# 5. Changes Made in Deployment Version
doc.add_heading('5. Changes Made in Deployment Version', level=1)
doc.add_paragraph('1. Extracted configurations to `config.yaml`.')
doc.add_paragraph('2. Added automated disk cleanup (retention policies).')
doc.add_paragraph('3. Implemented Rotating File Handlers for logs (10MB limit, keeping 5 backups).')
doc.add_paragraph('4. Created `install_service.bat` for easy NSSM service registration.')
doc.add_paragraph('5. Separated PaddleOCR into its own server (`paddleocr_server.py`) to fix PyTorch/PaddlePaddle CUDA/CPU conflicts.')
doc.add_paragraph('6. Optimized image processing thresholds (e.g., Laplacian variance) for better part presence detection.')

# 6. Top to Bottom Technical Details
doc.add_heading('6. Top to Bottom Technical Details', level=1)
doc.add_paragraph('Architecture Overview:')
arch = doc.add_paragraph(style='List Bullet')
doc.add_paragraph('Frontend: HTML/JS served via Flask on port 5000.')
doc.add_paragraph('Backend: Python 3.12, Flask, OpenCV (Image Processing), PyTorch/Ultralytics (YOLOv8 Segmentation).')
doc.add_paragraph('OCR Microservice: Flask server running on port 5001 strictly handling PaddleOCR inference.')
doc.add_paragraph('Hardware Interface: Uses OpenCV VideoCapture or IKapC SDK for industrial GigE cameras.')

doc.add_paragraph('Workflow:')
wf = doc.add_paragraph(style='List Number')
doc.add_paragraph('Step 1: Application starts, loads models (`SSB_FUEL_LID_V4.pt`), initializes camera stream.')
doc.add_paragraph('Step 2: Awaits the placement of a part (checked via HSV blue/metal masking).')
doc.add_paragraph('Step 3: Captures front side, detects defects, waits for flip.')
doc.add_paragraph('Step 4: Captures back side, crops serial number region, sends to PaddleOCR service.')
doc.add_paragraph('Step 5: Consolidates results, generates PDF report (`generate_report.py`), logs data, and resets for the next cycle.')

# Save Document
output_path = r'd:\SSP_FUEL_PRODUCTION\SSP_Fuel_Lid_Project_Handbook.docx'
doc.save(output_path)
print(f"Document saved successfully at {output_path}")
